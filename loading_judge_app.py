# -*- coding: utf-8 -*-
"""
loading_judge_app.py - 적재 판정 테스트 프로그램 (규칙 기반 + 딥러닝 하이브리드)

[사용 순서]
 1) OK 폴더 / NG 폴더 지정 (예: OK 50장, NG 30장)
 2) [ROI 지정] : 박스 '내부'(슬롯 줄무늬가 보이는 영역)를 드래그 후 Enter (최초 1회, 카메라 구도 바뀌면 재지정)
 3) [학습 시작] : 규칙 기준(각도/에너지) 자동 보정 + CNN(MobileNetV3-small) 전이학습
 4) [사진 업로드 및 판정] : 판정할 사진 여러 장 선택 -> 자동 판정
      저장 위치  : <저장폴더>/OK/<YYYY-MM-DD>/원본파일명.jpg
                   <저장폴더>/NG/<YYYY-MM-DD>/원본파일명.jpg
      엑셀       : <저장폴더>/judgement_<날짜_시간>.xlsx   (열: P Box Label QR | judgement)

[설치]
 pip install opencv-python numpy pillow openpyxl torch torchvision
 (CNN 학습 최초 1회는 사전학습 가중치 다운로드를 위해 인터넷 필요)
 torch 가 없어도 프로그램은 실행되며, 이 경우 규칙 기반만으로 판정합니다.
"""
import os
import sys
import io

# ------------------------------------------------------------------ PyInstaller windowed 모드 스트림 보호
class NullWriter:
    def write(self, s):
        pass
    def flush(self):
        pass

if sys.stdout is None:
    sys.stdout = NullWriter()
if sys.stderr is None:
    sys.stderr = NullWriter()

import json
import glob
import math
import queue
import shutil
import threading
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import numpy as np
import cv2
from PIL import Image, ImageTk
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ------------------------------------------------------------------ 설정
EXCEL_HEADERS = ["P Box Label QR", "judgement"]
STRIP_EXT = False            # True: 파일명에서 확장자(.jpg) 제거하여 엑셀에 기록
INPUT_SIZE = 224
MEAN = np.array([0.485, 0.456, 0.406], np.float32)
STD = np.array([0.229, 0.224, 0.225], np.float32)
IMG_EXTS = ("*.jpg", "*.jpeg", "*.png", "*.JPG", "*.JPEG", "*.PNG")


def get_base_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


BASE_DIR = get_base_dir()
MODEL_DIR = os.path.join(BASE_DIR, "models")
MODEL_PATH = os.path.join(MODEL_DIR, "loading_binary.pt")
CFG_PATH = os.path.join(BASE_DIR, "loading_judge_cfg.json")
DEFAULT_OUT = os.path.join(BASE_DIR, "judge_output")
os.makedirs(MODEL_DIR, exist_ok=True)

DEFAULT_CFG = {
    "roi": [0.38, 0.05, 0.72, 0.85],   # [x1,y1,x2,y2] 화면 비율 (임시값 - ROI 지정 버튼으로 재설정)
    "ref_theta_deg": None,
    "theta_tol_deg": 20.0,
    "min_edge_energy": 0.04,
    "min_coherence": 0.15,
    "cnn_min_prob": 0.90,              # CNN 의 OK 확률이 이 값 이상이어야 OK
}


# ------------------------------------------------------------------ 공통 유틸
def imread_unicode(path):
    try:
        return cv2.imdecode(np.fromfile(path, dtype=np.uint8), cv2.IMREAD_COLOR)
    except Exception:
        return None


def list_images(folder):
    files = []
    if folder and os.path.isdir(folder):
        for ext in IMG_EXTS:
            files += glob.glob(os.path.join(folder, ext))
    return sorted(set(files))


def load_cfg():
    cfg = dict(DEFAULT_CFG)
    if os.path.exists(CFG_PATH):
        try:
            with open(CFG_PATH, "r", encoding="utf-8") as f:
                cfg.update(json.load(f))
        except Exception:
            pass
    return cfg


def save_cfg(cfg):
    with open(CFG_PATH, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


def crop_roi(frame, roi):
    h, w = frame.shape[:2]
    x1, y1, x2, y2 = roi
    return frame[int(y1 * h):int(y2 * h), int(x1 * w):int(x2 * w)]


def ang_diff(a, b):
    return ((a - b + 90.0) % 180.0) - 90.0


def unique_path(path):
    """같은 이름이 이미 있으면 _1, _2 ... 를 붙여 덮어쓰기 방지"""
    if not os.path.exists(path):
        return path
    base, ext = os.path.splitext(path)
    n = 1
    while os.path.exists(f"{base}_{n}{ext}"):
        n += 1
    return f"{base}_{n}{ext}"


# ------------------------------------------------------------------ 규칙 기반
def rule_features(roi_bgr):
    """밝기가 아닌 '줄무늬 구조'(방향/정렬도/에지량)로 판단 (CLAHE + 구조 텐서)"""
    gray = cv2.cvtColor(roi_bgr, cv2.COLOR_BGR2GRAY)
    h, w = gray.shape[:2]
    s = 256.0 / max(h, w)
    gray = cv2.resize(gray, (max(8, int(w * s)), max(8, int(h * s))))
    gray = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(gray)
    gray = cv2.GaussianBlur(gray, (5, 5), 0)
    gx = cv2.Sobel(gray, cv2.CV_32F, 1, 0, ksize=3)
    gy = cv2.Sobel(gray, cv2.CV_32F, 0, 1, ksize=3)
    jxx, jyy, jxy = float(np.sum(gx * gx)), float(np.sum(gy * gy)), float(np.sum(gx * gy))
    tot = jxx + jyy + 1e-9
    return {
        "coherence": math.sqrt((jxx - jyy) ** 2 + 4 * jxy ** 2) / tot,
        "theta": 0.5 * math.degrees(math.atan2(2 * jxy, jxx - jyy)),
        "energy": float(np.mean(np.sqrt(gx * gx + gy * gy))) / 255.0,
    }


def rule_check(f, p):
    if f["energy"] < p["min_edge_energy"]:
        return False, "슬롯 패턴 미검출 (뚜껑 덮임/박스 없음 의심)"
    if f["coherence"] < p["min_coherence"]:
        return False, "슬롯 정렬 패턴 불명확"
    ref = p.get("ref_theta_deg")
    if ref is not None:
        d = ang_diff(f["theta"], ref)
        if abs(d) > p["theta_tol_deg"]:
            return False, f"적재 방향 불일치 (기준 대비 {d:+.0f}°) - 박스 회전 의심"
    return True, ""


def calibrate_rule(ok_files, roi):
    th, co, en = [], [], []
    for f in ok_files:
        img = imread_unicode(f)
        if img is None:
            continue
        ft = rule_features(crop_roi(img, roi))
        th.append(ft["theta"])
        co.append(ft["coherence"])
        en.append(ft["energy"])
    if not th:
        raise RuntimeError("OK 폴더에서 읽을 수 있는 이미지가 없습니다.")
    th, co, en = np.array(th), np.array(co), np.array(en)
    ref = math.degrees(np.angle(np.mean(np.exp(1j * np.radians(2 * th)))) / 2)
    diffs = np.abs([ang_diff(t, ref) for t in th])
    tol = float(np.clip(1.5 * np.percentile(diffs, 98), 10, 30))
    return {
        "ref_theta_deg": round(float(ref), 2),
        "theta_tol_deg": round(tol, 1),
        "min_edge_energy": round(float(0.6 * np.percentile(en, 5)), 4),
        "min_coherence": round(float(0.6 * np.percentile(co, 5)), 4),
    }


# ------------------------------------------------------------------ CNN 학습
def train_cnn(ok_files, ng_files, roi, log, epochs=25, batch=16):
    import torch
    import torch.nn as nn
    from torch.utils.data import Dataset, DataLoader
    from torchvision import models, transforms as T

    def split(files):
        files = sorted(files)               # 파일명 순 = 시간순 -> 뒤 20% 검증
        if len(files) < 5:
            return files, []
        k = max(1, int(len(files) * 0.2))
        return files[:-k], files[-k:]

    ok_tr, ok_va = split(ok_files)
    ng_tr, ng_va = split(ng_files)
    tr = [(f, 0) for f in ok_tr] + [(f, 1) for f in ng_tr]       # 0=OK, 1=NG
    va = [(f, 0) for f in ok_va] + [(f, 1) for f in ng_va]
    if not va:
        va = tr
        log("※ 사진이 적어 검증 세트를 따로 나누지 못했습니다 (학습 세트로 대체)")
    log(f"학습 {len(tr)}장 (OK {len(ok_tr)} / NG {len(ng_tr)}), 검증 {len(va)}장 (OK {len(ok_va)} / NG {len(ng_va)})")

    norm = T.Normalize(MEAN.tolist(), STD.tolist())
    tf_tr = T.Compose([
        T.ToPILImage(),
        T.ColorJitter(brightness=0.5, contrast=0.5, saturation=0.4, hue=0.05),
        T.RandomAffine(degrees=4, translate=(0.04, 0.04), scale=(0.95, 1.05)),
        T.RandomApply([T.GaussianBlur(5, (0.1, 1.5))], p=0.3),
        T.ToTensor(), norm,
        T.RandomErasing(p=0.3, scale=(0.02, 0.15)),
    ])
    tf_va = T.Compose([T.ToPILImage(), T.ToTensor(), norm])

    class DS(Dataset):
        def __init__(self, items, tf):
            self.items, self.tf = items, tf

        def __len__(self):
            return len(self.items)

        def __getitem__(self, i):
            path, y = self.items[i]
            img = imread_unicode(path)
            r = crop_roi(img, roi)
            r = cv2.cvtColor(cv2.resize(r, (INPUT_SIZE, INPUT_SIZE)), cv2.COLOR_BGR2RGB)
            return self.tf(r), y

    dev = "cuda" if torch.cuda.is_available() else "cpu"
    dl_tr = DataLoader(DS(tr, tf_tr), batch_size=batch, shuffle=True, num_workers=0)
    dl_va = DataLoader(DS(va, tf_va), batch_size=batch, shuffle=False, num_workers=0)

    try:
        # progress=False 로 진행률 바 터미널 출력 에러 방지
        weights = models.MobileNet_V3_Small_Weights.DEFAULT
        model = models.mobilenet_v3_small(weights=weights, progress=False)
    except Exception as e:
        raise RuntimeError(f"사전학습 가중치를 받지 못했습니다 (최초 1회 인터넷 필요): {e}")
        
    model.classifier[3] = nn.Linear(model.classifier[3].in_features, 2)
    model.to(dev)

    n_ok, n_ng = max(len(ok_tr), 1), max(len(ng_tr), 1)
    w = torch.tensor([(n_ok + n_ng) / (2 * n_ok), (n_ok + n_ng) / (2 * n_ng)], dtype=torch.float32).to(dev)
    crit = nn.CrossEntropyLoss(weight=w)
    opt = torch.optim.AdamW(model.parameters(), lr=3e-4, weight_decay=1e-4)
    sched = torch.optim.lr_scheduler.CosineAnnealingLR(opt, T_max=epochs)

    best, best_state, best_info = -9.0, None, ""
    for ep in range(1, epochs + 1):
        model.train()
        loss_sum = 0.0
        for x, y in dl_tr:
            x, y = x.to(dev), y.to(dev)
            opt.zero_grad()
            loss = crit(model(x), y)
            loss.backward()
            opt.step()
            loss_sum += loss.item() * len(x)
        sched.step()

        model.eval()
        correct = n = missed = 0
        with torch.no_grad():
            for x, y in dl_va:
                pred = model(x.to(dev)).argmax(1).cpu()
                correct += (pred == y).sum().item()
                n += len(y)
                missed += ((pred == 0) & (y == 1)).sum().item()     # NG를 OK로 오판
        acc = correct / max(n, 1)
        log(f"  [{ep:02d}/{epochs}] loss {loss_sum/max(len(tr),1):.4f} | 검증 정확도 {acc*100:.1f}% | NG→OK 오판 {missed}")
        score = acc - 0.5 * missed / max(n, 1)
        if score > best:
            best = score
            best_state = {k: v.cpu().clone() for k, v in model.state_dict().items()}
            best_info = f"검증 정확도 {acc*100:.1f}%, NG→OK 오판 {missed}건"

    torch.save(best_state, MODEL_PATH)
    return best_info


# ------------------------------------------------------------------ 판정기
class Judge:
    def __init__(self, cfg):
        self.cfg = cfg
        self.model = None
        self.torch = None
        self.load_model()

    def load_model(self):
        self.model = None
        if not os.path.exists(MODEL_PATH):
            return
        try:
            import torch
            import torch.nn as nn
            from torchvision import models
            m = models.mobilenet_v3_small(weights=None)
            m.classifier[3] = nn.Linear(m.classifier[3].in_features, 2)
            m.load_state_dict(torch.load(MODEL_PATH, map_location="cpu"))
            m.eval()
            self.model, self.torch = m, torch
        except Exception:
            self.model = None

    def _p_ok(self, roi_bgr):
        img = cv2.cvtColor(cv2.resize(roi_bgr, (INPUT_SIZE, INPUT_SIZE)), cv2.COLOR_BGR2RGB)
        img = (img.astype(np.float32) / 255.0 - MEAN) / STD
        x = self.torch.from_numpy(np.transpose(img, (2, 0, 1))[None].astype(np.float32))
        with self.torch.no_grad():
            prob = self.torch.softmax(self.model(x), 1)[0]
        return float(prob[0])

    def judge(self, img):
        p = self.cfg
        roi = crop_roi(img, p["roi"])
        if roi.size == 0:
            return {"verdict": "NG", "reason": "ROI 오류", "p_ok": None}
        rule_ok, rule_reason = rule_check(rule_features(roi), p)
        p_ok = self._p_ok(roi) if self.model is not None else None
        cnn_ok = True if p_ok is None else (p_ok >= p["cnn_min_prob"])

        ok = rule_ok and cnn_ok
        reason = ""
        if not ok:
            parts = []
            if not cnn_ok:
                parts.append(f"CNN 판정 NG (OK 확률 {p_ok*100:.0f}%)")
            if not rule_ok:
                parts.append(rule_reason)
            reason = " / ".join(parts)
        return {"verdict": "OK" if ok else "NG", "reason": reason, "p_ok": p_ok}


# ------------------------------------------------------------------ 엑셀
def write_excel(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "판정결과"
    ws.append(EXCEL_HEADERS)

    thin = Border(left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
                  top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))
    hfill = PatternFill(start_color="1F242D", end_color="1F242D", fill_type="solid")
    for c in range(1, 3):
        cell = ws.cell(row=1, column=c)
        cell.fill = hfill
        cell.font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin

    ok_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    ng_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    for name, verdict in rows:
        shown = os.path.splitext(name)[0] if STRIP_EXT else name
        ws.append([shown, verdict])
        r = ws.max_row
        for c in (1, 2):
            cell = ws.cell(row=r, column=c)
            cell.border = thin
            cell.alignment = Alignment(horizontal="left" if c == 1 else "center", vertical="center")
        ws.cell(row=r, column=2).fill = ok_fill if verdict == "OK" else ng_fill
        if verdict == "NG":
            ws.cell(row=r, column=2).font = Font(name="맑은 고딕", bold=True, color="C00000")

    width = max([len(str(r[0])) for r in rows] + [len(EXCEL_HEADERS[0])]) + 4
    ws.column_dimensions["A"].width = min(width, 90)
    ws.column_dimensions["B"].width = 14
    ws.freeze_panes = "A2"
    wb.save(path)


# ------------------------------------------------------------------ GUI
class App:
    def __init__(self, root):
        self.root = root
        root.title("적재 판정 테스트 프로그램 (규칙 기반 + 딥러닝)")
        root.geometry("1200x800")
        self.cfg = load_cfg()
        self.judge = Judge(self.cfg)
        self.q = queue.Queue()
        self.busy = False
        self.path_by_iid = {}
        self.photo = None
        self.cnt_ok = self.cnt_ng = 0

        self.var_ok = tk.StringVar()
        self.var_ng = tk.StringVar()
        self.var_out = tk.StringVar(value=DEFAULT_OUT)

        self.build()
        self.refresh_status()
        root.after(100, self.poll)

    def build(self):
        pad = {"padx": 8, "pady": 3}

        f1 = ttk.LabelFrame(self.root, text=" ① 학습 (초기 OK / NG 사진) ")
        f1.pack(fill=tk.X, padx=10, pady=(10, 4))
        for r, (label, var) in enumerate((("OK 사진 폴더", self.var_ok), ("NG 사진 폴더", self.var_ng))):
            ttk.Label(f1, text=label, width=12).grid(row=r, column=0, **pad)
            ttk.Entry(f1, textvariable=var, width=90).grid(row=r, column=1, **pad)
            ttk.Button(f1, text="찾기", command=lambda v=var: self.pick_dir(v)).grid(row=r, column=2, **pad)
        row3 = ttk.Frame(f1)
        row3.grid(row=2, column=0, columnspan=3, sticky="w", **pad)
        self.btn_roi = ttk.Button(row3, text="ROI 지정", command=self.on_roi)
        self.btn_roi.pack(side=tk.LEFT, padx=(0, 8))
        self.btn_train = ttk.Button(row3, text="학습 시작", command=self.on_train)
        self.btn_train.pack(side=tk.LEFT, padx=(0, 12))
        self.lbl_status = ttk.Label(row3, text="")
        self.lbl_status.pack(side=tk.LEFT)

        f2 = ttk.LabelFrame(self.root, text=" ② 판정 (사진 업로드) ")
        f2.pack(fill=tk.X, padx=10, pady=4)
        ttk.Label(f2, text="저장 폴더", width=12).grid(row=0, column=0, **pad)
        ttk.Entry(f2, textvariable=self.var_out, width=90).grid(row=0, column=1, **pad)
        ttk.Button(f2, text="찾기", command=lambda: self.pick_dir(self.var_out)).grid(row=0, column=2, **pad)
        row2 = ttk.Frame(f2)
        row2.grid(row=1, column=0, columnspan=3, sticky="w", **pad)
        self.btn_judge = ttk.Button(row2, text="사진 업로드 및 판정", command=self.on_judge)
        self.btn_judge.pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(row2, text="저장 폴더 열기", command=self.open_out).pack(side=tk.LEFT, padx=(0, 16))
        self.lbl_count = ttk.Label(row2, text="OK 0  /  NG 0", font=("맑은 고딕", 10, "bold"))
        self.lbl_count.pack(side=tk.LEFT)

        mid = ttk.Frame(self.root)
        mid.pack(fill=tk.BOTH, expand=True, padx=10, pady=4)

        cols = ("no", "name", "verdict", "pok", "reason")
        self.tree = ttk.Treeview(mid, columns=cols, show="headings", height=14)
        for c, t, w, a in (("no", "No", 45, "center"), ("name", "P Box Label QR (파일명)", 330, "w"),
                           ("verdict", "judgement", 80, "center"), ("pok", "OK확률", 70, "center"),
                           ("reason", "사유", 330, "w")):
            self.tree.heading(c, text=t)
            self.tree.column(c, width=w, anchor=a)
        self.tree.tag_configure("ng", background="#fde2e2", foreground="#b00020")
        sb = ttk.Scrollbar(mid, orient=tk.VERTICAL, command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        sb.pack(side=tk.LEFT, fill=tk.Y)
        self.tree.bind("<<TreeviewSelect>>", self.on_select)

        self.lbl_prev = ttk.Label(mid, text="(행을 선택하면 사진과 ROI 영역이 표시됩니다)", anchor="center")
        self.lbl_prev.pack(side=tk.LEFT, fill=tk.BOTH, padx=(8, 0))

        f3 = ttk.LabelFrame(self.root, text=" 로그 ")
        f3.pack(fill=tk.BOTH, padx=10, pady=(4, 10))
        self.txt = tk.Text(f3, height=10, font=("Consolas", 9))
        self.txt.pack(fill=tk.BOTH, expand=True)

    def log(self, msg):
        self.q.put(("log", msg))

    def pick_dir(self, var):
        d = filedialog.askdirectory(initialdir=var.get() or BASE_DIR)
        if d:
            var.set(d)

    def open_out(self):
        d = self.var_out.get().strip() or DEFAULT_OUT
        os.makedirs(d, exist_ok=True)
        try:
            os.startfile(d)
        except Exception:
            messagebox.showinfo("저장 폴더", d)

    def set_busy(self, b):
        self.busy = b
        st = "disabled" if b else "normal"
        for btn in (self.btn_roi, self.btn_train, self.btn_judge):
            btn.config(state=st)

    def refresh_status(self):
        c = self.cfg
        model = "CNN 모델 있음" if self.judge.model is not None else "CNN 모델 없음(규칙 기반만 사용)"
        cal = "규칙 보정 완료" if c.get("ref_theta_deg") is not None else "규칙 미보정"
        self.lbl_status.config(text=f"ROI {c['roi']}  |  {cal}  |  {model}")

    def poll(self):
        try:
            while True:
                kind, data = self.q.get_nowait()
                if kind == "log":
                    self.txt.insert(tk.END, data + "\n")
                    self.txt.see(tk.END)
                elif kind == "busy":
                    self.set_busy(data)
                    self.refresh_status()
                elif kind == "row":
                    i, name, res, path = data
                    p = "-" if res["p_ok"] is None else f"{res['p_ok']*100:.0f}%"
                    iid = self.tree.insert("", tk.END, values=(i, name, res["verdict"], p, res["reason"]),
                                           tags=("ng",) if res["verdict"] == "NG" else ())
                    self.path_by_iid[iid] = path
                    self.tree.see(iid)
                    if res["verdict"] == "OK":
                        self.cnt_ok += 1
                    else:
                        self.cnt_ng += 1
                    self.lbl_count.config(text=f"OK {self.cnt_ok}  /  NG {self.cnt_ng}")
        except queue.Empty:
            pass
        self.root.after(100, self.poll)

    def on_roi(self):
        path = filedialog.askopenfilename(
            title="ROI 지정용 샘플 사진 선택 (OK 사진 권장)", initialdir=self.var_ok.get() or BASE_DIR,
            filetypes=[("Images", "*.jpg *.jpeg *.png")])
        if not path:
            return
        img = imread_unicode(path)
        if img is None:
            messagebox.showerror("오류", "이미지를 읽을 수 없습니다.")
            return
        h, w = img.shape[:2]
        messagebox.showinfo("ROI 지정", "박스 '내부'(슬롯 줄무늬가 보이는 영역)를 드래그한 뒤 Enter 를 누르세요.\n"
                                       "박스 테두리와 바깥 물체는 제외할수록 정확합니다.")
        x, y, rw, rh = cv2.selectROI("ROI (drag, then Enter)", img, showCrosshair=False)
        cv2.destroyAllWindows()
        if rw == 0 or rh == 0:
            return
        self.cfg["roi"] = [round(x / w, 4), round(y / h, 4), round((x + rw) / w, 4), round((y + rh) / h, 4)]
        self.cfg["ref_theta_deg"] = None
        save_cfg(self.cfg)
        self.refresh_status()
        self.log(f"ROI 저장: {self.cfg['roi']}  -> [학습 시작]을 다시 실행하세요.")

    def on_train(self):
        ok_files = list_images(self.var_ok.get().strip())
        ng_files = list_images(self.var_ng.get().strip())
        if not ok_files or not ng_files:
            messagebox.showwarning("폴더 확인", "OK 폴더와 NG 폴더에 이미지가 있어야 합니다.")
            return
        self.set_busy(True)
        threading.Thread(target=self._train_worker, args=(ok_files, ng_files), daemon=True).start()

    def _train_worker(self, ok_files, ng_files):
        try:
            roi = self.cfg["roi"]
            self.log(f"== 학습 시작: OK {len(ok_files)}장 / NG {len(ng_files)}장 ==")

            cal = calibrate_rule(ok_files, roi)
            self.cfg.update(cal)
            save_cfg(self.cfg)
            self.log(f"[규칙 기반 보정] {cal}")

            ok_pass = sum(1 for f in ok_files
                          if (im := imread_unicode(f)) is not None
                          and rule_check(rule_features(crop_roi(im, roi)), self.cfg)[0])
            ng_catch = sum(1 for f in ng_files
                           if (im := imread_unicode(f)) is not None
                           and not rule_check(rule_features(crop_roi(im, roi)), self.cfg)[0])
            self.log(f"[규칙 단독 성능] OK 통과 {ok_pass}/{len(ok_files)}, NG 검출 {ng_catch}/{len(ng_files)}")

            self.log("[CNN 학습] 시작 (PC 사양에 따라 수 분 소요)")
            info = train_cnn(ok_files, ng_files, roi, self.log)
            self.judge.cfg = self.cfg
            self.judge.load_model()
            self.log(f"[CNN 학습 완료] 최적 모델: {info}")
            self.log("== 학습 완료. 이제 [사진 업로드 및 판정]으로 학습에 쓰지 않은 사진을 판정해 보세요 ==")
        except ImportError:
            self.log("※ torch / torchvision 이 설치되어 있지 않습니다. 규칙 기반 보정만 완료되었습니다.")
            self.log("   CNN 사용: pip install torch torchvision")
        except Exception as e:
            self.log(f"오류: {e}")
        finally:
            self.q.put(("busy", False))

    def on_judge(self):
        files = filedialog.askopenfilenames(title="판정할 사진 선택 (여러 장 가능)",
                                            filetypes=[("Images", "*.jpg *.jpeg *.png")])
        if not files:
            return
        if self.cfg.get("ref_theta_deg") is None:
            if not messagebox.askyesno("학습 필요", "아직 규칙 보정(학습)이 되어 있지 않습니다.\n"
                                                 "그래도 현재 기본값으로 판정할까요?"):
                return
        out_root = self.var_out.get().strip() or DEFAULT_OUT
        self.tree.delete(*self.tree.get_children())
        self.path_by_iid.clear()
        self.cnt_ok = self.cnt_ng = 0
        self.lbl_count.config(text="OK 0  /  NG 0")
        self.set_busy(True)
        threading.Thread(target=self._judge_worker, args=(list(files), out_root), daemon=True).start()

    def _judge_worker(self, files, out_root):
        try:
            now = datetime.now()
            day = now.strftime("%Y-%m-%d")
            rows = []
            if self.judge.model is None:
                self.log("※ CNN 모델이 없어 규칙 기반만으로 판정합니다.")
            self.log(f"== 판정 시작: {len(files)}장 ==")
            for i, f in enumerate(files, 1):
                name = os.path.basename(f)
                img = imread_unicode(f)
                if img is None:
                    res = {"verdict": "NG", "reason": "이미지 읽기 실패", "p_ok": None}
                else:
                    res = self.judge.judge(img)

                dst_dir = os.path.join(out_root, res["verdict"], day)
                os.makedirs(dst_dir, exist_ok=True)
                try:
                    shutil.copy2(f, unique_path(os.path.join(dst_dir, name)))
                except Exception as e:
                    self.log(f"저장 실패 {name}: {e}")

                rows.append((name, res["verdict"]))
                self.q.put(("row", (i, name, res, f)))

            os.makedirs(out_root, exist_ok=True)
            xlsx = os.path.join(out_root, f"judgement_{now.strftime('%Y-%m-%d_%H%M%S')}.xlsx")
            write_excel(xlsx, rows)
            n_ok = sum(1 for _, v in rows if v == "OK")
            self.log(f"== 판정 완료: OK {n_ok} / NG {len(rows) - n_ok} ==")
            self.log(f"사진 저장: {out_root}\\{'OK|NG'}\\{day}\\  |  엑셀: {xlsx}")
        except Exception as e:
            self.log(f"오류: {e}")
        finally:
            self.q.put(("busy", False))

    def on_select(self, event=None):
        sel = self.tree.selection()
        if not sel:
            return
        path = self.path_by_iid.get(sel[0])
        img = imread_unicode(path) if path else None
        if img is None:
            return
        h, w = img.shape[:2]
        x1, y1, x2, y2 = self.cfg["roi"]
        cv2.rectangle(img, (int(x1 * w), int(y1 * h)), (int(x2 * w), int(y2 * h)), (0, 255, 255), 3)
        s = 460.0 / w
        img = cv2.resize(img, (460, int(h * s)))
        self.photo = ImageTk.PhotoImage(Image.fromarray(cv2.cvtColor(img, cv2.COLOR_BGR2RGB)))
        self.lbl_prev.config(image=self.photo, text="")


if __name__ == "__main__":
    root = tk.Tk()
    App(root)
    root.mainloop()
